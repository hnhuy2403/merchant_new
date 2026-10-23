#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
Tests for xlsx-utils.py — exercises inspect, write-testcase, and read-cases
against the user's KBKTCN template if available, with a synthetic fallback
template otherwise.

Run:
    python -m pytest scripts/tests/test_xlsx_utils.py -v

Or directly:
    python scripts/tests/test_xlsx_utils.py
"""

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).parent.parent
SCRIPT_PATH = SCRIPT_DIR / "xlsx-utils.py"

# The user's actual template (if present)
PROJECT_ROOT = SCRIPT_DIR.parent.parent.parent
USER_TEMPLATE = PROJECT_ROOT / "docs" / "noti-autorecharge-endofday-phase2" / "KBKTCN_template_testcase_yennth72_v1.0.xlsx"


def make_synthetic_template(path: Path) -> None:
    """Create a synthetic template that mirrors the KBKTCN structure."""
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Trang bìa"
    cover["A1"] = "Cover page"

    # Add a 'template' sheet matching the KBKTCN layout
    ws = wb.create_sheet("template")
    ws["B1"] = "KỊCH BẢN KIỂM THỬ CHỨC NĂNG"
    ws["B3"] = "Tên màn hình/Tên chức năng"
    ws["C3"] = ""  # filled by write-testcase
    ws["B4"] = "Mã trường hợp kiểm thử"
    ws["C4"] = ""
    ws["B6"] = "Manual test"
    ws["B7"] = "Số trường hợp kiểm thử đạt (P)"
    ws["C7"] = '=COUNTIF($N$16:$N$1702,"P")'
    ws["B8"] = "Số trường hợp kiểm thử không đạt (F)"
    ws["C8"] = '=COUNTIF($N$16:$N$387,"F")'
    ws["B11"] = "Tổng số trường hợp kiểm thử"
    ws["C11"] = '=COUNTA($F$16:$F$382)'

    # Header row 13
    ws["A13"] = "ID"
    ws["B13"] = "Mục đích kiểm thử"
    ws["C13"] = "Trường hợp kiểm thử"
    ws["D13"] = "Data test"
    ws["E13"] = "Các bước thực hiện"
    ws["F13"] = "Kết quả mong muốn"
    ws["G13"] = "Ảnh màn hình"
    ws["N13"] = "Kết quả (M)"
    ws["O13"] = "Mã lỗi"
    ws["P13"] = "Ghi chú"

    # Row 14 sub-headers
    ws["H14"] = "IOS"
    ws["I14"] = "Android"

    wb.save(path)


def run(*args, capture=True):
    """Run xlsx-utils.py with given args, return (returncode, stdout)."""
    cmd = [sys.executable, str(SCRIPT_PATH), *args]
    result = subprocess.run(cmd, capture_output=capture, text=True, encoding="utf-8")
    return result.returncode, result.stdout, result.stderr


def get_template() -> Path:
    """Return path to user's template if it exists, else create synthetic."""
    if USER_TEMPLATE.exists():
        return USER_TEMPLATE
    tmp = Path(tempfile.gettempdir()) / "qc_synthetic_template.xlsx"
    make_synthetic_template(tmp)
    return tmp


def test_inspect():
    template = get_template()
    rc, out, err = run("inspect", str(template))
    assert rc == 0, f"inspect failed: {err}"
    data = json.loads(out)
    assert "sheets" in data
    assert "template_sheet" in data
    assert data["template_sheet"].lower() == "template"
    assert "ID" in data["expected_columns"]
    print("[OK] test_inspect")


def test_write_and_read_roundtrip():
    template = get_template()
    cases = [
        {"suite": "Suite-1: Login validation", "id": "AR-001",
         "purpose": "Verify accountId rejects letters",
         "case": "Type letters in accountId",
         "data": "abcXYZ",
         "steps": "1. Open Login\n2. Type abcXYZ in accountId\n3. Tap Continue",
         "expected": "Field rejects, error message shows",
         "note": ""},
        {"suite": "Suite-1: Login validation", "id": "AR-002",
         "purpose": "Verify accountId accepts digits only",
         "case": "Type digits in accountId",
         "data": "1234567890",
         "steps": "1. Open Login\n2. Type 1234567890 in accountId\n3. Tap Continue",
         "expected": "Field accepts, screen advances",
         "note": ""},
        {"suite": "Suite-2: Boundary checks", "id": "AR-003",
         "purpose": "Verify accountId max length",
         "case": "Type 19-digit accountId",
         "data": "1" * 19,
         "steps": "1. Open Login\n2. Type 1*19 in accountId\n3. Tap Continue",
         "expected": "Field truncates to 18 digits",
         "note": "Boundary"},
    ]

    with tempfile.TemporaryDirectory() as td:
        cases_path = Path(td) / "cases.json"
        cases_path.write_text(json.dumps(cases, ensure_ascii=False), encoding="utf-8")
        output_path = Path(td) / "output.xlsx"

        rc, out, err = run(
            "write-testcase",
            "--template", str(template),
            "--output", str(output_path),
            "--feature", "Auto Recharge - Login validation",
            "--id-prefix", "AR",
            "--cases-json", str(cases_path),
            "--sheet-name", "AutoRecharge_Login",
        )
        assert rc == 0, f"write-testcase failed: {err}"
        result = json.loads(out)
        assert result["status"] == "success"
        assert result["case_count"] == 3
        assert result["suite_count"] == 2
        assert output_path.exists()

        # Verify the output file: open it and check content
        wb = openpyxl.load_workbook(output_path, data_only=False)
        assert "AutoRecharge_Login" in wb.sheetnames
        ws = wb["AutoRecharge_Login"]

        # R3 col 3 should be the feature name
        assert ws.cell(row=3, column=3).value == "Auto Recharge - Login validation"
        # R4 col 3 should be the ID prefix
        assert ws.cell(row=4, column=3).value == "AR"
        # R7 should still have a COUNTIF formula (preserved from template)
        formula = ws.cell(row=7, column=3).value
        assert formula and isinstance(formula, str) and formula.startswith("=COUNTIF"), \
            f"Expected COUNTIF formula in R7C3, got: {formula!r}"

        # Now read cases back and verify
        rc, out, err = run("read-cases", str(output_path), "--sheet", "AutoRecharge_Login")
        assert rc == 0, f"read-cases failed: {err}"
        read_result = json.loads(out)
        assert read_result["case_count"] == 3
        ids = [c["id"] for c in read_result["cases"]]
        assert ids == ["AR-001", "AR-002", "AR-003"]
        # First two cases should have suite "Suite-1: Login validation"
        assert read_result["cases"][0]["suite"].startswith("Suite-1")
        assert read_result["cases"][2]["suite"].startswith("Suite-2")

    print("[OK] test_write_and_read_roundtrip")


def test_missing_template():
    rc, out, err = run("write-testcase",
                       "--template", "/nonexistent/template.xlsx",
                       "--output", "/tmp/x.xlsx",
                       "--feature", "X",
                       "--id-prefix", "X",
                       "--cases-json", "/nonexistent/cases.json")
    assert rc == 1, "expected non-zero exit on missing template"
    assert "not found" in err
    print("[OK] test_missing_template")


def main() -> int:
    tests = [test_inspect, test_write_and_read_roundtrip, test_missing_template]
    failed = 0
    for t in tests:
        try:
            t()
        except AssertionError as e:
            print(f"[FAIL] {t.__name__}: {e}")
            failed += 1
        except Exception as e:
            print(f"[ERROR] {t.__name__}: {e}")
            failed += 1
    print(f"\n{len(tests) - failed}/{len(tests)} passed.")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
